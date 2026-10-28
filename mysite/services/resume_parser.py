from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_RESUME_BYTES = 15 * 1024 * 1024
MAX_RESUME_CHARS = 25_000
MAX_PDF_PAGES = 100
MIN_RESUME_CHARS = 50


class ResumeParseError(ValueError):
    pass


def extract_resume_text(filename: str, data: bytes) -> str:
    if not data:
        raise ResumeParseError("Файл пустой.")
    if len(data) > MAX_RESUME_BYTES:
        raise ResumeParseError("Файл превышает ограничение 15 МБ.")

    extension = Path(filename).suffix.lower()
    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise ResumeParseError("PDF защищён паролем. Пришлите незашифрованный файл.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ResumeParseError("В PDF слишком много страниц (максимум 100).")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            document = Document(BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            paragraphs.extend(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
            text = "\n".join(paragraphs)
        elif extension == ".txt":
            text = data.decode("utf-8-sig")
        else:
            raise ResumeParseError("Поддерживаются только PDF, DOCX и TXT.")
    except ResumeParseError:
        raise
    except (PdfReadError, BadZipFile, OSError, ValueError, KeyError) as exc:
        raise ResumeParseError("Не удалось прочитать файл. Проверьте, что он не повреждён.") from exc

    text = "\n".join(line.strip() for line in text.replace("\x00", "").splitlines()).strip()
    if len(text) < MIN_RESUME_CHARS:
        if extension == ".pdf":
            raise ResumeParseError(
                "В PDF не найден читаемый текст. Возможно, это скан — распознайте его с помощью OCR."
            )
        raise ResumeParseError("В документе слишком мало читаемого текста (нужно не менее 50 символов).")
    return text[:MAX_RESUME_CHARS]
